import csv
import io
import json
from datetime import datetime, timezone as dt_timezone
from zoneinfo import ZoneInfo

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.urls import reverse
from django.views.decorators.http import require_GET, require_POST
from openlexiconApp.models import Lang

from . import conf, services
from .filters import FILTERS, build_checks, presets_for
from .generator import (
    ALGORITHMS,
    MAX_LENGTH,
    MIN_LENGTH,
    GenerationError,
    allowed_algorithms,
    default_algorithm,
    generate_pseudowords,
)

MAX_DOWNLOAD_COLS = 20


def _error(message, status=400):
    return JsonResponse({"error": message}, status=status)


def _json_body(request):
    try:
        data = json.loads(request.body.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        return None
    return data if isinstance(data, dict) else None


def _language_payload(language):
    """Database used for a language + default constraints for that language."""
    database = services.get_database_for_language(language)
    payload = {"language": language, "presets": presets_for(language), "database": None, "gram_classes": []}
    if database is not None:
        payload["database"] = {
            "name": database.name,
            "info": database.info or "",
            "website": database.website or "",
        }
        payload["gram_classes"] = services.get_gram_classes(database)
    return payload


# --------------------------------------------------------------------------- pages


@require_GET
def index(request):
    languages = services.list_languages()
    primary_languages =  [x for x in languages if x in [Lang.FR, Lang.EN, Lang.DE]]
    secondary_languages = [x for x in languages if x not in primary_languages]
    default_language = conf.get("DEFAULT_LANGUAGE")
    config = {
        "urls": {
            "language": reverse("unipseudo:language"),
            "words": reverse("unipseudo:words"),
            "generate": reverse("unipseudo:generate"),
            "download": reverse("unipseudo:download"),
        },
        "minLength": MIN_LENGTH,
        "maxLength": MAX_LENGTH,
        "minPseudos": conf.get("MIN_PSEUDOWORDS"),
        "maxPseudos": conf.get("MAX_PSEUDOWORDS"),
        "algorithms": {
            str(n): {"allowed": allowed_algorithms(n), "default": default_algorithm(n)}
            for n in range(MIN_LENGTH, MAX_LENGTH + 1)
        },
        "filters": [f.name for f in FILTERS],
    }
    return render(
        request,
        "unipseudo/index.html",
        {
            "languages": languages,
            "primary_languages": primary_languages,
            "secondary_languages": secondary_languages,
            "default_language": default_language if default_language in languages else "",
            "filters": FILTERS,
            "lengths": range(MIN_LENGTH, MAX_LENGTH + 1),
            "default_length": 6,
            "default_pseudowords": conf.get("DEFAULT_PSEUDOWORDS"),
            "min_pseudowords": conf.get("MIN_PSEUDOWORDS"),
            "max_pseudowords": conf.get("MAX_PSEUDOWORDS"),
            "config": config,
        },
    )


# --------------------------------------------------------------------------- API


@require_GET
def language_detail(request):
    """Database used for a language, its grammatical classes, default constraints."""
    return JsonResponse(_language_payload(request.GET.get("language", "")))


def _parse_length(raw):
    try:
        length = int(raw)
    except (TypeError, ValueError):
        raise GenerationError("Invalid pseudoword length.")
    if not MIN_LENGTH <= length <= MAX_LENGTH:
        raise GenerationError("Pseudoword length must be between %d and %d." % (MIN_LENGTH, MAX_LENGTH))
    return length


def _get_database(language):
    """Database for a language; None for "Other" / unknown language."""
    return services.get_database_for_language(language) if isinstance(language, str) else None


@require_GET
def words(request):
    """'Use words from database': models of the right length (and class)."""
    try:
        length = _parse_length(request.GET.get("length"))
        database = _get_database(request.GET.get("language"))
    except GenerationError as exc:
        return _error(str(exc))
    if database is None:
        return JsonResponse({"words": []})
    try:
        found = services.get_words(database, length, gram_class=request.GET.get("gram_class") or None)
    except services.FrequencyColumnError as exc:
        return _error(str(exc))
    return JsonResponse({"words": found})


@require_POST
def generate(request):
    data = _json_body(request)
    if data is None:
        return _error("Invalid request.")

    min_n, max_n = conf.get("MIN_PSEUDOWORDS"), conf.get("MAX_PSEUDOWORDS")
    try:
        length = _parse_length(data.get("length"))
        database = _get_database(data.get("language"))
        algorithm = data.get("algorithm")
        if algorithm not in ALGORITHMS:
            raise GenerationError("Wrong substring type selected.")
        if algorithm not in allowed_algorithms(length):
            raise GenerationError("The %s algorithm cannot be used for %d-letter pseudowords." % (algorithm, length))
        try:
            n = int(data.get("n"))
        except (TypeError, ValueError):
            n = None
        if n is None or not min_n <= n <= max_n:
            raise GenerationError("Please enter a Number of pseudowords between %d and %d." % (min_n, max_n))
    except GenerationError as exc:
        return _error(str(exc))

    # Source words: one per line, lower-cased, right length, no duplicates
    raw_words = str(data.get("words") or "").splitlines()
    models = list(dict.fromkeys(w.strip().lower() for w in raw_words if len(w.strip()) == length))
    if len(models) > conf.get("MAX_MODEL_WORDS"):
        return _error("Too many source words.")

    exclude = services.get_words(database, length, need_full_list=True) if database else ()
    checks = build_checks(data.get("filters"))

    try:
        result = generate_pseudowords(
            n=n, length=length, models=models, algorithm=algorithm, exclude=exclude, checks=checks
        )
    except GenerationError as exc:
        if str(exc) == "no_models":
            return _error('Please provide words of the desired length in the "Words to use" section !')
        return _error(str(exc))

    found = len(result.pseudowords)
    if result.complete:
        info = "%d pseudowords generated from %d source words !" % (found, result.n_models)
        warning = ""
    else:
        info = (
            "%d pseudowords generated from %d source words..." % (found, result.n_models)
            if found
            else "Failed to generate pseudowords from %d source words..." % result.n_models
        )
        warning = (
            "Failed to generate the requested number of pseudowords: not enough valid source words provided. "
            "We generated %d of the %d requested. Please enter words of the desired length in the "
            '"Words to use" section. The more pseudowords you need, the more words you need to provide. '
            "A minimum of 100 words for 20 pseudowords is a good basis, unfortunately there is no fixed minimal "
            'number of words. You need to proceed with trial-and-error or use the "Word Import Tool" for easier '
            "pseudowords generation." % (found, n)
        )

    return JsonResponse(
        {
            "rows": [{"p": p.pseudoword, "s": p.sources} for p in result.pseudowords],
            "subLen": result.sub_len,
            "info": info,
            "warning": warning,
        }
    )


def _clean_cell(value):
    return value if isinstance(value, str) else ""


@require_POST
def download(request):
    """Excel / TSV export of the rows currently shown in the table."""
    data = _json_body(request)
    if data is None:
        return _error("Invalid request.")

    rows = data.get("rows")
    if not isinstance(rows, list) or not rows or len(rows) > conf.get("MAX_PSEUDOWORDS"):
        return _error("Nothing to download.")
    table = []
    for row in rows:
        if not isinstance(row, list) or not 0 < len(row) <= MAX_DOWNLOAD_COLS:
            return _error("Invalid rows.")
        table.append([_clean_cell(c) for c in row])

    width = max(len(r) for r in table)
    detail = bool(data.get("detail"))
    header = ["Pseudoword"] + ["Word.%d" % i for i in range(1, width)]
    table = [r + [""] * (width - len(r)) for r in table]

    try:
        now = datetime.now(ZoneInfo(str(data.get("tz") or "UTC")))
    except Exception:
        now = datetime.now(dt_timezone.utc)
    stamp = now.strftime("%Y-%m-%dT%H-%M-%S%z")
    base = "Pseudowords-details-query-" if detail else "Pseudowords-query-"

    if data.get("format") == "tsv":
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        writer.writerows(table)
        response = HttpResponse(buf.getvalue(), content_type="text/tab-separated-values; charset=utf-8")
        filename = "%s%s.tsv" % (base, stamp)
    else:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Pseudowords"
        ws.append(header)
        for row in table:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                    cell.data_type = "s"  # never interpret user text as a formula
        out = io.BytesIO()
        wb.save(out)
        response = HttpResponse(
            out.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename = "%s%s.xlsx" % (base, stamp)
    response["Content-Disposition"] = 'attachment; filename="%s"' % filename
    return response
